#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""从台账 + 候选池自动生成三份专项报告的骨架与全部数字。

为什么必须有这一步（开箱即用的最后一个断点）
------------------------------------------
实测：`run_pipeline` 跑到台账就停了，三份报告要 agent 从零发明。
而 SKILL.md 把三份报告列为「交付主体，不是附录」——
「一条命令跑通」和「报告要自己写」是矛盾的。

本脚本不假装能写出好文章，它做的是**把难且易错的部分自动化**：
  · 所有数字从台账/候选池复算，绝不手抄
  · 遍历性：逐信源/逐来源的覆盖统计
  · 真实性：链接可核性、置信度分布、缺口清单
  · 匹配性：判定链落点分布、逐条判据
  · 留出「人写」的位置（标注 TODO），不替人编造分析

用法
----
  python scripts/make_reports.py --ledger 台账.xlsx --candidates 候选.jsonl \
      --outdir ./报告 --month 2026-09
  python scripts/make_reports.py --ledger 台账.xlsx --outdir ./报告 --month 2026-09 \
      --excluded 排除.json --manual 需人工判.json

退出码：0 = 三份报告已生成 | 2 = 台账为空 | 1 = 无法判定
"""
import argparse
import json
import os
import re
import sys
from collections import Counter

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass

try:
    import openpyxl
except Exception as e:  # noqa: BLE001
    print("无法导入 openpyxl：%s\n处置： python -m pip install openpyxl" % e)
    sys.exit(1)


def read_ledger(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [("" if c.value is None else str(c.value).strip()) for c in ws[1]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n) + 1
        return None

    C = {
        "摘要": col("事件摘要"), "国家": col("国家"), "日期": col("日期"),
        "年份": col("年份"), "月份": col("月份"), "地点": col("地点"),
        "死亡": col("死亡人数"), "受伤": col("受伤人数"),
        "方式": col("袭击方式"), "组织": col("负责人/组织", "袭击者"),
        "URL": col("信息来源URL"), "置信度": col("置信度"),
        "状态": col("审核状态"), "知识点": col("关联知识点"), "备注": col("备注"),
    }

    def num(v):
        try:
            return int(float(str(v).replace(",", "")))
        except Exception:  # noqa: BLE001
            return 0

    rows = []
    for r in range(2, (ws.max_row or 1) + 1):
        if all(ws.cell(r, c).value is None for c in range(1, (ws.max_column or 1) + 1)):
            continue
        first = str(ws.cell(r, 1).value or "")
        if re.match(r"^(合计|总计|小计|total)", first, re.I):
            continue
        def g(k):
            c = C.get(k)
            return ws.cell(r, c).value if c else None
        if not str(g("国家") or "").strip():
            continue
        rows.append({
            "row": r,
            "summary": str(g("摘要") or ""),
            "country": str(g("国家") or "").strip(),
            "year": g("年份"), "month": g("月份"), "day": g("日期"),
            "place": str(g("地点") or ""),
            "deaths": num(g("死亡")), "injuries": num(g("受伤")),
            "act": str(g("方式") or ""),
            "org": str(g("组织") or ""),
            "url": str(g("URL") or "").strip(),
            "conf": str(g("置信度") or ""),
            "status": str(g("状态") or ""),
            "chain": str(g("知识点") or ""),
            "note": str(g("备注") or ""),
        })
    return rows


def read_jsonl(path):
    if not path or not os.path.isfile(path):
        return []
    out = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    out.append(json.loads(line))
                except Exception:  # noqa: BLE001
                    pass
    return out


def read_json(path):
    if not path or not os.path.isfile(path):
        return []
    try:
        d = json.load(open(path, encoding="utf-8"))
        return d if isinstance(d, list) else []
    except Exception:  # noqa: BLE001
        return []


def fmt(n):
    return "{:,}".format(n) if isinstance(n, int) else str(n)


# ---------------------------------------------------------------- 报告一：遍历性
def report_coverage(rows, cand, month, outdir):
    lines = []
    A = lines.append
    A("# %s 事件台账 · 遍历性报告" % month)
    A("")
    A("> **这份报告回答一个问题：该看的信源，看了没有。**")
    A("> 数字全部由 `scripts/make_reports.py` 从台账与候选池复算，未手工填写。")
    A("")
    A("## 一、规模")
    A("")
    A("| 项 | 值 |")
    A("|---|---|")
    A("| 目标时段 | %s |" % month)
    A("| 候选池条数 | %s |" % fmt(len(cand)))
    A("| 计入台账 | %s |" % fmt(len(rows)))
    A("| 覆盖国家/地区 | %s |" % fmt(len({r["country"] for r in rows})))
    A("")
    src = Counter(c.get("source") or "(未标注)" for c in cand)
    if src:
        A("## 二、数据来源构成")
        A("")
        A("| 来源 | 候选条数 |")
        A("|---|---|")
        for k, v in src.most_common():
            A("| %s | %s |" % (k, fmt(v)))
        A("")
    A("## 三、国别覆盖")
    A("")
    A("| 国家/地区 | 计入条数 |")
    A("|---|---|")
    for k, v in Counter(r["country"] for r in rows).most_common():
        A("| %s | %d |" % (k, v))
    A("")
    A("## 四、必须交底的缺口")
    A("")
    A("<!-- TODO(人写)：逐条填下列四项，不许省略 -->")
    A("")
    A("1. **未覆盖的国家/时段**：<哪些国家本月 0 条，为什么>")
    A("2. **数据源上限**：<所用源本身收什么、不收什么>")
    A("3. **未检索的部分**：<哪些信源/通道没试过>")
    A("4. **渠道不通**：<哪些源取不到内容，与「没报道」区分开>")
    A("")
    A("## 五、这份报告不能证明什么")
    A("")
    A("- 不能证明该月「只有这些事件」—— 只能证明「在已检索范围内确证了这些」。")
    A("- 不能证明没有遗漏 —— 见第四节的缺口清单。")
    A("")
    p = os.path.join(outdir, "报告-1-遍历性.md")
    open(p, "w", encoding="utf-8").write("\n".join(lines))
    return p


# ---------------------------------------------------------------- 报告二：真实性
def report_authenticity(rows, cand, month, outdir):
    lines = []
    A = lines.append
    A("# %s 事件台账 · 真实性报告" % month)
    A("")
    A("> **这份报告回答一个问题：这些事件是真的吗，有没有虚构。**")
    A("> 数字由 `make_reports.py` 复算。")
    A("")
    n = len(rows)
    with_url = [r for r in rows if r["url"].startswith("http")]
    no_url = [r for r in rows if not r["url"].startswith("http")]
    A("## 一、总体")
    A("")
    A("| 指标 | 值 |")
    A("|---|---|")
    A("| 计入事件 | %d |" % n)
    A("| **有真实 URL 的** | **%d（%.0f%%）** |" % (len(with_url), 100.0 * len(with_url) / n if n else 0))
    A("| 缺 URL / URL 非链接 | %d |" % len(no_url))
    A("| 死亡合计 | %s |" % fmt(sum(r["deaths"] for r in rows)))
    A("| 受伤合计 | %s |" % fmt(sum(r["injuries"] for r in rows)))
    A("")
    if no_url:
        A("## 二、缺可核链接的条目（**必须补齐或降级**）")
        A("")
        A("| 台账行 | 国家 | 日期 | 摘要 |")
        A("|---|---|---|---|")
        for r in no_url[:40]:
            A("| %d | %s | %s-%s-%s | %s |"
              % (r["row"], r["country"], r["year"], r["month"], r["day"], r["summary"][:50]))
        if len(no_url) > 40:
            A("| ... | 另有 %d 条 | | |" % (len(no_url) - 40))
        A("")
        A("> **处置**：这些行必须补上真实出版商链接，否则按口径应为「无法确证」不计入。")
        A("")
    A("## 三、置信度分布")
    A("")
    A("| 档 | 条数 |")
    A("|---|---|")
    for k, v in Counter(r["conf"] or "(未填)" for r in rows).most_common():
        A("| %s | %d |" % (k, v))
    A("")
    A("## 四、数字口径说明")
    A("")
    A("<!-- TODO(人写)：说明伤亡数字的取值规则与滚动更新情况 -->")
    A("")
    A("- 死亡/受伤按「有具体数字取数、数十=50/十数=15/数=5、提及无数字记 1、未提记 0」。")
    A("- 若同一事件多源数字不同，取最终值并在备注写区间。")
    A("")
    A("## 五、这份报告不能证明什么")
    A("")
    A("- 没有做逐条人工事实核查（未向当地媒体二次比对）。")
    A("- 单源条目不得宣称「多源确证」。")
    A("")
    p = os.path.join(outdir, "报告-2-真实性.md")
    open(p, "w", encoding="utf-8").write("\n".join(lines))
    return p


# ---------------------------------------------------------------- 报告三：匹配性
def report_fit(rows, excluded, manual, month, outdir):
    lines = []
    A = lines.append
    A("# %s 事件台账 · 匹配性报告" % month)
    A("")
    A("> **这份报告回答一个问题：这些事件符合口径定义吗。**")
    A("> 判定链落点分布由 `make_reports.py` 复算。")
    A("")
    A("## 一、判定结果")
    A("")
    A("| 落点 | 条数 |")
    A("|---|---|")
    A("| **计入** | **%d** |" % len(rows))
    A("| 不计入 | %d |" % len(excluded))
    A("| 需人工判 | %d |" % len(manual))
    A("")
    chains = Counter()
    for r in rows:
        c = (r["chain"] or "").strip() or "(未标)"
        chains[c] += 1
    if chains:
        A("## 二、计入事件命中的判据")
        A("")
        A("| 判据 | 条数 |")
        A("|---|---|")
        for k, v in chains.most_common():
            A("| %s | %d |" % (k, v))
        A("")
    exc_reason = Counter()
    for e in excluded:
        note = str(e.get("备注") or e.get("note") or "")
        m = re.search(r"\[(?:排除|判定)\]\s*([^\s，,。]{1,20})", note)
        exc_reason[m.group(1) if m else "(未标)"] += 1
    if exc_reason:
        A("## 三、被排除的落点分布")
        A("")
        A("| 落点 | 条数 |")
        A("|---|---|")
        for k, v in exc_reason.most_common():
            A("| %s | %d |" % (k, v))
        A("")
    if manual:
        A("## 四、需人工判的条目（%d 条）" % len(manual))
        A("")
        A("> 「需人工判」不是「待定」，是明确的动作：逐条看原文判完写回。")
        A("")
        A("| 行 | 国家 | 日期 | 摘要 |")
        A("|---|---|---|---|")
        for m in manual[:30]:
            A("| %s | %s | %s | %s |"
              % (m.get("row", ""), m.get("国家", ""),
                 m.get("日期", ""), str(m.get("事件摘要", ""))[:46]))
        if len(manual) > 30:
            A("| ... | 另有 %d 条 | | |" % (len(manual) - 30))
        A("")
    A("## 五、争议与边界")
    A("")
    A("<!-- TODO(人写)：列出哪些判定最可能被推翻、为什么 -->")
    A("")
    A("- 口径张力：标准没写清的情形，我们怎么裁的")
    A("- 争议条目：多源数字不一致、主体存疑的")
    A("- 已知误判风险：哪几条最可能被推翻")
    A("")
    p = os.path.join(outdir, "报告-3-匹配性.md")
    open(p, "w", encoding="utf-8").write("\n".join(lines))
    return p


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ledger", required=True)
    ap.add_argument("--candidates", default=None, help="候选池 jsonl")
    ap.add_argument("--excluded", default=None, help="被排除的 json")
    ap.add_argument("--manual", default=None, help="需人工判的 json")
    ap.add_argument("--month", required=True)
    ap.add_argument("--outdir", default="./报告")
    a = ap.parse_args()

    if not os.path.isfile(a.ledger):
        print("台账不存在：%s" % a.ledger)
        return 1
    rows = read_ledger(a.ledger)
    if not rows:
        print("台账里没有数据行，无法生成报告。")
        return 2

    os.makedirs(a.outdir, exist_ok=True)
    cand = read_jsonl(a.candidates)
    exc = read_json(a.excluded)
    man = read_json(a.manual)

    p1 = report_coverage(rows, cand, a.month, a.outdir)
    p2 = report_authenticity(rows, cand, a.month, a.outdir)
    p3 = report_fit(rows, exc, man, a.month, a.outdir)

    print("=" * 70)
    print(" 三份专项报告已生成")
    print("=" * 70)
    print("台账条数 : %d" % len(rows))
    print("候选池   : %d 条" % len(cand))
    print("排除     : %d 条" % len(exc))
    print("需人工判 : %d 条" % len(man))
    print()
    for p in (p1, p2, p3):
        print("  → %s" % p)
    print()
    print("⚠️  报告里有 TODO 标记的段落必须人写（缺口清单、争议边界、数字口径）。")
    print("    脚本只保证数字准确，不替你写分析。")
    print("=" * 70)
    return 0


if __name__ == "__main__":
    sys.exit(main())
